"""Geometry + text-fit audit for generated PPTX decks.

Checks:
  1. off-slide shapes
  2. node-vs-node bounding box overlaps
  3. connector lines passing through the INTERIOR of a node
     (zero-thickness vertical/horizontal lines get an epsilon so they
      are not silently skipped -- this was missed in earlier versions)
  4. text overflow: wrap-aware height for rects, per-line width for diamonds
"""
import sys
import re
import math
from pptx import Presentation

EMU = 914400
CHAR_W_FACTOR = 0.52
LINE_H_FACTOR = 1.15
EPS = 0.02  # treat a zero-thickness line as this thick for overlap tests

NODE_GEOMS = {"roundRect", "diamond", "ellipse", "rect"}


def to_in(v):
    return v / EMU


def prst(shp):
    m = re.search(r'<a:prstGeom prst="([a-zA-Z0-9]+)"', shp._element.xml)
    return m.group(1) if m else None


def max_font(shp):
    sizes = re.findall(r'sz="(\d+)"', shp._element.xml)
    return max(int(s) for s in sizes) / 100 if sizes else None


def audit(path):
    p = Presentation(path)
    sw, sh = to_in(p.slide_width), to_in(p.slide_height)
    issues = 0

    for si, slide in enumerate(p.slides, 1):
        nodes, lines = [], []
        for shp in slide.shapes:
            if shp.left is None:
                continue
            l, t, w, h = to_in(shp.left), to_in(shp.top), to_in(shp.width), to_in(shp.height)
            geom = prst(shp)
            # A table's frame height is a placeholder -- PowerPoint grows the frame
            # to the sum of its row heights. Measure the real extent, or tables
            # silently overlap whatever is placed below them.
            if shp.has_table:
                h = sum(to_in(r.height) for r in shp.table.rows)
                geom = "table"
                first = shp.table.cell(0, 0).text.strip().replace("\n", " ")
                nodes.append((l, t, w, h, "table", f"[table] {first[:44]}", shp))
                if l < -0.02 or t < -0.02 or l + w > sw + 0.02 or t + h > sh + 0.02:
                    print(f"S{si} OFF-SLIDE table [{l:.2f},{t:.2f},{w:.2f},{h:.2f}]")
                    issues += 1
                continue
            text = shp.text_frame.text.replace("\n", " / ")[:52] if shp.has_text_frame else ""
            if l < -0.02 or t < -0.02 or l + w > sw + 0.02 or t + h > sh + 0.02:
                print(f"S{si} OFF-SLIDE geom={geom} [{l:.2f},{t:.2f},{w:.2f},{h:.2f}] \"{text}\"")
                issues += 1
            if geom == "line":
                lines.append((l, t, w, h))
            elif geom in NODE_GEOMS and (shp.has_text_frame or w > 0.2):
                nodes.append((l, t, w, h, geom, text, shp))

        # solid boxes only (plain text boxes legitimately overlay cards/chips)
        solid = [n for n in nodes if n[4] != "rect"]
        tables = [n for n in nodes if n[4] == "table"]
        # a table overlapping ANY visible element is a defect, rects included
        for tb in tables:
            for other in nodes:
                if other is tb:
                    continue
                l1, t1, w1, h1 = tb[0], tb[1], tb[2], tb[3]
                l2, t2, w2, h2 = other[0], other[1], other[2], other[3]
                ox = max(0, min(l1 + w1, l2 + w2) - max(l1, l2))
                oy = max(0, min(t1 + h1, t2 + h2) - max(t1, t2))
                if ox > 0.05 and oy > 0.05:
                    print(f"S{si} TABLE-OVERLAP ({ox:.2f}x{oy:.2f}) \"{tb[5]}\" <-> \"{other[5]}\"")
                    issues += 1

        # 2. node vs node
        for i in range(len(solid)):
            for j in range(i + 1, len(solid)):
                l1, t1, w1, h1, g1, x1, _ = solid[i]
                l2, t2, w2, h2, g2, x2, _ = solid[j]
                ox = max(0, min(l1 + w1, l2 + w2) - max(l1, l2))
                oy = max(0, min(t1 + h1, t2 + h2) - max(t1, t2))
                if ox > 0.02 and oy > 0.02:
                    print(f"S{si} NODE-OVERLAP ({ox:.2f}x{oy:.2f}) \"{x1}\" <-> \"{x2}\"")
                    issues += 1

        # 2b. content straddling a card edge. A filled roundRect is a container:
        # any text placed on top of it should be wholly inside. Partly-in /
        # partly-out is the signature of a list that has outgrown its slide.
        cards = [n for n in nodes if n[4] == "roundRect" and n[2] > 1.5 and n[3] > 0.5]
        for (cl, ct, cw, ch, _, ctext, _) in cards:
            for (nl, nt, nw, nh, ngeom, ntext, _) in nodes:
                if ngeom != "rect" or not ntext:
                    continue
                ox = max(0, min(cl + cw, nl + nw) - max(cl, nl))
                oy = max(0, min(ct + ch, nt + nh) - max(ct, nt))
                if ox <= 0.05 or oy <= 0.05:
                    continue
                inside = (nl >= cl - 0.05 and nt >= ct - 0.05
                          and nl + nw <= cl + cw + 0.05 and nt + nh <= ct + ch + 0.05)
                if not inside:
                    print(f"S{si} STRADDLES-CARD \"{ntext}\" hangs over card \"{ctext}\"")
                    issues += 1

        # 3. line through node interior (epsilon-padded so axis-aligned lines count)
        for (ll, lt, lw, lh) in lines:
            pl, pt = ll, lt
            pw, ph = max(lw, EPS), max(lh, EPS)
            for (nl, nt, nw, nh, geom, text, _) in nodes:
                ox = max(0, min(pl + pw, nl + nw) - max(pl, nl))
                oy = max(0, min(pt + ph, nt + nh) - max(pt, nt))
                if ox <= 0.01 or oy <= 0.01:
                    continue
                ocx = max(pl, nl) + ox / 2
                ocy = max(pt, nt) + oy / 2
                if (nl + nw * 0.22 < ocx < nl + nw * 0.78) and (nt + nh * 0.22 < ocy < nt + nh * 0.78):
                    print(f"S{si} ARROW-THROUGH-NODE ({ox:.2f}x{oy:.2f}) node=\"{text}\"")
                    issues += 1

        # 4. text fit
        for l, t, w, h, geom, text, shp in nodes:
            if geom not in ("roundRect", "diamond"):
                continue
            fs = max_font(shp)
            if fs is None:
                continue
            if geom == "roundRect":
                cap = (w * 0.90) / (fs / 72 * CHAR_W_FACTOR)
                line_h = fs / 72 * LINE_H_FACTOR
                total = 0
                for para in shp.text_frame.paragraphs:
                    txt = para.text
                    total += max(1, math.ceil(len(txt) / cap)) if txt else 1
                need = total * line_h + 0.07
                if need > h + 0.01:
                    print(f"S{si} RECT-OVERFLOW box=({w:.2f}x{h:.2f}) need={need:.2f} fs={fs} \"{text}\"")
                    issues += 1
            else:
                cap = (w * 0.50) / (fs / 72 * CHAR_W_FACTOR)
                paras = [pa.text for pa in shp.text_frame.paragraphs]
                for txt in paras:
                    if txt and len(txt) > cap:
                        print(f"S{si} DIAMOND-LINE-TOO-LONG \"{txt}\" ({len(txt)}>{cap:.0f})")
                        issues += 1
                need_v = len(paras) * (fs / 72 * LINE_H_FACTOR)
                if need_v > h * 0.62:
                    print(f"S{si} DIAMOND-TOO-TALL need={need_v:.2f} usable={h*0.62:.2f} \"{text}\"")
                    issues += 1

    print(f"\nTOTAL ISSUES: {issues}")
    return issues


if __name__ == "__main__":
    sys.exit(1 if audit(sys.argv[1]) else 0)
